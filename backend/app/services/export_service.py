"""统一导出服务（ADR 000 §P-3：WPS 兼容 docx/xlsx；导出失败降级 md/纯文本复制）
依赖懒加载（函数内 import），缺依赖不影响应用启动。
"""
import uuid

from app.config import settings


def _new_filename(prefix: str, ext: str) -> str:
    return f"{prefix}_{uuid.uuid4().hex[:8]}{ext}"


def _set_cn(run, name: str = "宋体"):
    """设置中文字体（eastAsia），WPS/Office 双兼容。"""
    run.font.name = name  # 同时写 w:ascii / w:hAnsi
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is not None:
        rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", name)


def export_docx(title: str, blocks: list[dict]) -> str:
    """python-docx 生成 docx，返回文件名（相对 EXPORT_DIR）。
    blocks 块类型：{"type":"para","text"} / {"type":"list","items"} / {"type":"table","headers","rows"}
    """
    from docx import Document
    from docx.shared import Pt, Cm

    doc = Document()
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(3.18)

    h = doc.add_heading(title, level=1)
    for run in h.runs:
        _set_cn(run, "黑体")
        run.font.size = Pt(18)

    for block in blocks or []:
        btype = block.get("type", "para")
        if btype == "para":
            p = doc.add_paragraph(str(block.get("text", "")))
            for run in p.runs:
                _set_cn(run)
        elif btype == "list":
            for item in block.get("items", []):
                p = doc.add_paragraph(str(item), style="List Bullet")
                for run in p.runs:
                    _set_cn(run)
        elif btype == "table":
            headers = block.get("headers", []) or []
            rows = block.get("rows", []) or []
            table = doc.add_table(rows=1 + len(rows), cols=max(len(headers), 1))
            table.style = "Table Grid"
            for j, htext in enumerate(headers):
                cell = table.cell(0, j)
                cell.text = str(htext)
                for run in cell.paragraphs[0].runs:
                    _set_cn(run, "黑体")
            for i, row in enumerate(rows, start=1):
                for j, value in enumerate(row):
                    cell = table.cell(i, j)
                    cell.text = str(value)
                    for run in cell.paragraphs[0].runs:
                        _set_cn(run)

    filename = _new_filename("export", ".docx")
    settings.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(str(settings.EXPORT_DIR / filename))
    return filename


def export_xlsx(headers: list, rows: list[list]) -> str:
    """openpyxl 生成 xlsx。金额列请传打码值（amount_masked）。返回文件名。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    if headers:
        ws.append([str(h) for h in headers])
        for cell in ws[1]:
            cell.font = Font(bold=True)
    for row in rows:
        ws.append([str(v) if v is not None else "" for v in row])

    filename = _new_filename("export", ".xlsx")
    settings.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(str(settings.EXPORT_DIR / filename))
    return filename
